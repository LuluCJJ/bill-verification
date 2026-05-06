from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[3]
OUT = ROOT / "artifacts" / "business-deck" / "output"
PREVIEWS = ROOT / "artifacts" / "business-deck" / "previews"
OUT.mkdir(parents=True, exist_ok=True)
PREVIEWS.mkdir(parents=True, exist_ok=True)

W, H = 1600, 900
BG = "#F5F7FA"
INK = "#162333"
MUTED = "#5B6878"
BLUE = "#1F66B3"
GREEN = "#227C55"
RED = "#B93A3A"
AMBER = "#B87514"
LINE = "#D7E0EA"

FONT = "C:/Windows/Fonts/msyh.ttc"
FONT_BOLD = "C:/Windows/Fonts/msyhbd.ttc"


def font(size, bold=False):
    return ImageFont.truetype(FONT_BOLD if bold else FONT, size)


def draw_text(draw, xy, text, size=32, fill=INK, bold=False, max_width=None, line_gap=8):
    x, y = xy
    f = font(size, bold)
    if "\n" in text:
        for line in text.split("\n"):
            y = draw_text(draw, (x, y), line, size, fill, bold, max_width, line_gap)
        return y
    if not max_width:
        draw.text((x, y), text, font=f, fill=fill)
        return y + draw.textbbox((x, y), text, font=f)[3] - y
    lines = []
    cur = ""
    for ch in text:
        trial = cur + ch
        if draw.textlength(trial, font=f) <= max_width or not cur:
            cur = trial
        else:
            lines.append(cur)
            cur = ch
    if cur:
        lines.append(cur)
    for line in lines:
        draw.text((x, y), line, font=f, fill=fill)
        y += size + line_gap
    return y


def rounded(draw, box, fill, outline=None, width=2, r=22):
    draw.rounded_rectangle(box, radius=r, fill=fill, outline=outline, width=width)


def pill(draw, x, y, text, color=BLUE):
    f = font(24, True)
    w = draw.textlength(text, font=f) + 44
    rounded(draw, (x, y, x + w, y + 48), fill="#EAF3FF", outline="#BFD7F0", r=24)
    draw.text((x + 22, y + 12), text, font=f, fill=color)
    return x + w


def arrow(draw, x1, y1, x2, y2, color=BLUE):
    draw.line((x1, y1, x2, y2), fill=color, width=5)
    draw.polygon([(x2, y2), (x2 - 18, y2 - 10), (x2 - 18, y2 + 10)], fill=color)


def slide_canvas():
    return Image.new("RGB", (W, H), BG)


def footer(draw, idx):
    draw.text((80, 842), f"Bill Verification AI Pre-audit · {idx}/7", font=font(18), fill="#7A8795")


def save_slide(img, idx):
    path = PREVIEWS / f"slide_{idx:02d}.png"
    img.save(path)
    return path


slides = []


def s1():
    img = slide_canvas()
    d = ImageDraw.Draw(img)
    d.rectangle((0, 0, W, H), fill="#EDF4F8")
    d.polygon([(980, 0), (1600, 0), (1600, 900), (1220, 900)], fill="#D7EAF4")
    d.polygon([(1120, 0), (1600, 0), (1600, 900), (1360, 900)], fill="#C8E0EA")
    draw_text(d, (90, 110), "权签票据一致性\nAI 预审方案", 72, INK, True, 760, 16)
    draw_text(d, (94, 320), "从“逐项肉眼核对”走向“AI 风险初筛 + 权签人聚焦审查”", 34, MUTED, False, 760)
    x = 94
    for label in ["支票", "转账信", "汇款申请书", "MAP/AP 指令"]:
        x = pill(d, x, 430, label) + 14
    rounded(d, (1040, 190, 1450, 620), "#FFFFFF", "#BFD0DD", 3, 28)
    draw_text(d, (1090, 250), "系统指令", 30, BLUE, True)
    draw_text(d, (1090, 310), "金额 USD 12,500\n收款方 Global Supplier\n账号 987654321000", 26, INK, False, 300, 14)
    arrow(d, 1130, 500, 1340, 500)
    draw_text(d, (1110, 540), "AI 预审", 28, GREEN, True)
    footer(d, 1)
    return img


def s2():
    img = slide_canvas()
    d = ImageDraw.Draw(img)
    draw_text(d, (80, 62), "为什么要做：权签核对正在变成高成本风险点", 48, INK, True)
    cards = [
        ("票面形态多", "全球 200+ 银行模板，字段叫法、语言、版式差异大。", BLUE),
        ("系统字段多", "MAP/AP 有几十个结构化字段，票面往往只出现子集。", AMBER),
        ("人工压力大", "600+ 权签人，兼职权签多，靠肉眼逐项核对体验差。", RED),
    ]
    x = 80
    for title, body, color in cards:
        rounded(d, (x, 190, x + 440, 510), "#FFFFFF", LINE, 2, 22)
        d.ellipse((x + 28, 220, x + 88, 280), fill=color)
        draw_text(d, (x + 110, 222), title, 34, INK, True)
        draw_text(d, (x + 36, 320), body, 28, MUTED, False, 360, 12)
        x += 500
    draw_text(d, (100, 650), "目标不是替代权签人，而是在权签前做一轮风险初筛，让人聚焦真正需要判断的问题。", 38, GREEN, True, 1320)
    footer(d, 2)
    return img


def s3():
    img = slide_canvas()
    d = ImageDraw.Draw(img)
    draw_text(d, (80, 62), "主流程：一笔付款任务如何完成 AI 预审", 48, INK, True)
    steps = [
        ("1", "系统指令", "MAP/AP 带出付款字段"),
        ("2", "票据文件", "上传或查看扫描件"),
        ("3", "模板匹配", "系统带出 + AI 辅助"),
        ("4", "AI 提取", "识别票面字段并归一"),
        ("5", "风险核验", "只突出不一致风险"),
        ("6", "人工反馈", "沉淀为调优样本"),
    ]
    xs = [80, 320, 560, 800, 1040, 1280]
    for i, (n, title, body) in enumerate(steps):
        x = xs[i]
        d.ellipse((x, 240, x + 120, 360), fill="#EAF3FF", outline="#BFD7F0", width=3)
        draw_text(d, (x + 43, 270), n, 42, BLUE, True)
        draw_text(d, (x - 10, 405), title, 30, INK, True, 150)
        draw_text(d, (x - 32, 455), body, 22, MUTED, False, 190)
        if i < len(steps) - 1:
            arrow(d, x + 132, 300, xs[i + 1] - 22, 300, "#8BA8C5")
    rounded(d, (180, 660, 1420, 760), "#FFFFFF", "#CFE0D6", 2, 26)
    draw_text(d, (230, 692), "Demo 中的“开始 AI 预审”会真实调用多模态模型，提取字段后自动进入规则核验。", 32, GREEN, True, 1120)
    footer(d, 3)
    return img


def s4():
    img = slide_canvas()
    d = ImageDraw.Draw(img)
    draw_text(d, (80, 62), "产品页面：四类用户动作分开，不再混在一个调试台", 46, INK, True, 1320)
    items = [
        ("付款核验", "权签人", "看系统指令、票面文件、AI 提取和风险结果", BLUE),
        ("模板调优", "配置人员", "维护模板字段叫法、位置提示和 AI 识别要求", GREEN),
        ("反馈样本", "运营/AI 产品", "聚合 AI 识别错误，转成可复用规则", AMBER),
        ("系统设置", "管理员", "配置模型 URL、模型名、API Key 和连通测试", "#64748B"),
    ]
    y = 160
    for title, role, body, color in items:
        rounded(d, (110, y, 1490, y + 120), "#FFFFFF", LINE, 2, 20)
        d.rectangle((110, y, 128, y + 120), fill=color)
        draw_text(d, (160, y + 25), title, 34, INK, True)
        pill(d, 405, y + 30, role, color)
        draw_text(d, (650, y + 32), body, 28, MUTED, False, 760)
        y += 145
    footer(d, 4)
    return img


def s5():
    img = slide_canvas()
    d = ImageDraw.Draw(img)
    draw_text(d, (80, 62), "核验结果：默认只让权签人看到风险", 48, INK, True)
    rounded(d, (90, 170, 710, 670), "#FFFFFF", LINE, 2, 22)
    rounded(d, (890, 170, 1510, 670), "#FFFFFF", LINE, 2, 22)
    draw_text(d, (130, 215), "系统支付指令", 34, BLUE, True)
    draw_text(d, (930, 215), "票面 AI 提取值", 34, GREEN, True)
    rows = [("收款方", "Global Supplier Limited", "Global Supplier Limited", True), ("账号", "987654321000", "987654321000", True), ("金额", "12,500.00", "15,200.00", False), ("币种", "USD", "USD", True)]
    y = 300
    for label, sys, doc, ok in rows:
        draw_text(d, (130, y), label, 24, MUTED, True)
        draw_text(d, (250, y), sys, 26, INK)
        draw_text(d, (930, y), label, 24, MUTED, True)
        if not ok:
            rounded(d, (1000, y - 12, 1420, y + 52), "#FFF1F1", "#E3B0B0", 2, 14)
        draw_text(d, (1050, y), doc, 26, RED if not ok else INK, True if not ok else False)
        if not ok:
            draw_text(d, (1220, y + 82), "高风险：金额不一致", 32, RED, True)
        y += 90
    draw_text(d, (130, 735), "通过项默认收起；风险项显性展示系统值、票面值和差异原因。", 34, INK, True, 1180)
    footer(d, 5)
    return img


def s6():
    img = slide_canvas()
    d = ImageDraw.Draw(img)
    draw_text(d, (80, 62), "AI 调优闭环：不是改代码，而是把业务经验变成模板规则", 44, INK, True, 1360)
    centers = [(210, 350), (520, 350), (830, 350), (1140, 350), (1370, 350)]
    labels = [
        ("AI 识别", "模型读票面"),
        ("人工纠错", "权签反馈"),
        ("样本沉淀", "进入反馈池"),
        ("模板调优", "别名位置要求"),
        ("下次改善", "进入提示词"),
    ]
    for i, ((cx, cy), (title, sub)) in enumerate(zip(centers, labels)):
        d.ellipse((cx - 85, cy - 85, cx + 85, cy + 85), fill="#FFFFFF", outline="#BFD0DD", width=3)
        draw_text(d, (cx - 62, cy - 26), title, 28, INK, True, 130)
        draw_text(d, (cx - 56, cy + 28), sub, 20, MUTED, False, 120)
        if i < len(centers) - 1:
            arrow(d, cx + 95, cy, centers[i + 1][0] - 95, cy, "#8BA8C5")
    rounded(d, (180, 620, 1420, 750), "#FFFFFF", "#CFE0D6", 2, 24)
    draw_text(d, (230, 650), "例：Intermediary Bank 不能映射为收款方银行 → 添加到该银行模板的 AI 识别要求 → 下一次提取时进入 prompt。", 32, GREEN, True, 1120)
    footer(d, 6)
    return img


def s7():
    img = slide_canvas()
    d = ImageDraw.Draw(img)
    draw_text(d, (80, 62), "当前已验证什么，下一步补什么", 48, INK, True)
    rounded(d, (90, 170, 730, 690), "#FFFFFF", "#CFE0D6", 2, 22)
    rounded(d, (870, 170, 1510, 690), "#FFFFFF", "#F0D7A8", 2, 22)
    draw_text(d, (130, 220), "当前 Demo 已跑通", 36, GREEN, True)
    done = ["火山方舟 Doubao 多模态接入", "三组合成样例端到端验证", "AI 提取 → 规则核验 → 风险展示", "模板调优配置进入 prompt"]
    y = 300
    for item in done:
        draw_text(d, (135, y), "✓", 30, GREEN, True)
        draw_text(d, (180, y), item, 28, INK, False, 480)
        y += 72
    draw_text(d, (910, 220), "下一轮建议", 36, AMBER, True)
    todo = ["模板识别：系统带出 + AI 建议 + 人工修正", "反馈样本池：聚合高频错误并转规则", "PDF 上传：转图片后进入模型识别", "更贴近 MAP 的任务和字段口径"]
    y = 300
    for item in todo:
        draw_text(d, (915, y), "→", 30, AMBER, True)
        draw_text(d, (960, y), item, 28, INK, False, 480)
        y += 72
    footer(d, 7)
    return img


slides = [s1(), s2(), s3(), s4(), s5(), s6(), s7()]
preview_paths = [save_slide(img, i + 1) for i, img in enumerate(slides)]


def add_textbox(slide, x, y, w, h, text_value, size=24, bold=False, color=INK):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text_value
    p.font.name = "Microsoft YaHei"
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = RGBColor.from_string(color.strip("#"))
    return box


def add_rect(slide, x, y, w, h, color, line=LINE, radius=False):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = RGBColor.from_string(color.strip("#"))
    shp.line.color.rgb = RGBColor.from_string(line.strip("#"))
    return shp


prs = Presentation()
prs.slide_width = Inches(16)
prs.slide_height = Inches(9)
blank = prs.slide_layouts[6]
for idx, png in enumerate(preview_paths, start=1):
    slide = prs.slides.add_slide(blank)
    slide.shapes.add_picture(str(png), 0, 0, width=Inches(16), height=Inches(9))
    # Add a small editable title overlay off the main visual hierarchy so the PPTX remains searchable/editable.
    add_textbox(slide, 0.15, 8.55, 6, 0.3, f"权签票据一致性 AI 预审 - 第 {idx} 页", 8, False, "#7A8795")

pptx_path = OUT / "bill_verification_business_overview.pptx"
prs.save(pptx_path)

montage = Image.new("RGB", (W * 2, H * 4), "#FFFFFF")
for i, img in enumerate(slides):
    thumb = img.resize((W, H))
    montage.paste(thumb, ((i % 2) * W, (i // 2) * H))
montage_path = PREVIEWS / "montage.png"
montage.save(montage_path)

print(pptx_path)
for p in preview_paths:
    print(p)
print(montage_path)
